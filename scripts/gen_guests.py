#!/usr/bin/env python3
"""Generate guests.xlsx — the master guest list for the wedding invitation.

Re-run any time:  uv run --with openpyxl python3 scripts/gen_guests.py
Edit the GUESTS list below to manage your guest list.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

SITE_URL = "https://chwbyte.github.io/wedding-invitation-dummy/"

# (Nama, Kategori, Status, Jumlah Tamu, WhatsApp, Catatan)
# Status: "Belum Diundang" | "Menunggu Jawaban" | "Hadir" | "Berhalangan"
GUESTS = [
    ("Andi Saputra", "Keluarga", "Menunggu Jawaban", 2, "+628111111111", "Sepupu dari pihak mempelai pria"),
    ("Bunga Melati", "Teman Kuliah", "Menunggu Jawaban", 1, "+628222222222", "Sahabat mempelai wanita"),
    ("Keluarga Hartono", "Keluarga", "Menunggu Jawaban", 4, "+628333333333", "Paman & bibi"),
    ("Rizky Pratama", "Teman Kerja", "Menunggu Jawaban", 1, "+628444444444", ""),
    ("Siti Rahayu", "Keluarga", "Menunggu Jawaban", 3, "+628555555555", "Kakek & nenek"),
]

HEADERS = ["No", "Nama", "Link Undangan", "Kategori", "Status", "Jumlah Tamu", "WhatsApp", "Catatan"]

GOLD = "B8975A"
GOLD_SOFT = "E7D9C0"
INK = "3A322C"

thin = Side(style="thin", color="D8C9AC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_fill = PatternFill("solid", fgColor=GOLD)
header_font = Font(bold=True, color="FFFFFF", size=11)
alt_fill = PatternFill("solid", fgColor="F7F1E6")


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Guests"

    # Header
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    # Rows
    for i, g in enumerate(GUESTS, start=1):
        name, category, status, guests, wa, note = g
        link = SITE_URL + "?name=" + name.replace(" ", "+")
        row = [i, name, link, category, status, guests, wa, note]
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=i + 1, column=c, value=val)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=(c == 8))
            if c == 2:
                cell.font = Font(bold=True, color=INK)
            if c == 3:
                cell.hyperlink = link
                cell.font = Font(color="0563C1", underline="single")
            if c in (5, 6):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if i % 2 == 0:
                cell.fill = alt_fill

    # Column widths
    widths = [5, 22, 60, 16, 18, 13, 16, 34]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{len(GUESTS) + 1}"

    # Status dropdown
    dv = DataValidation(
        type="list",
        formula1='"Belum Diundang,Menunggu Jawaban,Hadir,Berhalangan"',
        allow_blank=True,
    )
    ws.add_data_validation(dv)
    dv.add(f"E2:E{len(GUESTS) + 200}")

    # --- Second sheet: how-to ---
    ws2 = wb.create_sheet("Petunjuk")
    lines = [
        ["Cara memakai file ini (Guest List)"],
        [""],
        ["1. Kolom 'Link Undangan' sudah berisi link personal untuk setiap tamu."],
        ["   Kirim link itu (WhatsApp/email) — nama tamu otomatis muncul di undangan."],
        ["2. Ganti kolom 'Status' saat tamu menjawab: Menunggu Jawaban -> Hadir / Berhalangan."],
        ["3. Untuk menambah tamu: isi baris baru, salin pola link dari baris di atasnya."],
        ["4. RSVP yang masuk lewat website otomatis tercatat di Google Sheet terpisah"],
        ["   (lihat README.md, bagian 'RSVP & Wishes -> Google Sheets')."],
        ["5. File ini adalah salinan offline (master list). Live responses ada di Google Sheet."],
    ]
    for r, row in enumerate(lines, 1):
        cell = ws2.cell(row=r, column=1, value=row[0])
        if r == 1:
            cell.font = Font(bold=True, size=13, color=INK)
        elif row[0].startswith(("1.", "2.", "3.", "4.", "5.")):
            cell.font = Font(bold=True)
        ws2.column_dimensions["A"].width = 90

    out = "guests.xlsx"
    wb.save(out)
    print(f"Wrote {out} with {len(GUESTS)} guests")


if __name__ == "__main__":
    build()
